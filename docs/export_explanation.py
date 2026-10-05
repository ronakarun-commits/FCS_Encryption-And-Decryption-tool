"""Export the project code explanation to Word and PDF."""

from pathlib import Path

try:
    from docx import Document
    from docx.shared import Inches, Pt
except ModuleNotFoundError:  # Word export remains optional when the package is unavailable.
    Document = None
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, Preformatted, SimpleDocTemplate, Spacer


DOCS_DIR = Path(__file__).parent
SOURCE = DOCS_DIR / "project_code_explanation.md"
WORD_OUTPUT = DOCS_DIR / "project_code_explanation.docx"
PDF_OUTPUT = DOCS_DIR / "project_code_explanation.pdf"


def markdown_lines():
    return SOURCE.read_text(encoding="utf-8").splitlines()


def build_word(lines):
    if Document is None:
        return False

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    for line in lines:
        if not line.strip():
            continue
        if line.startswith("# "):
            document.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("```"):
            continue
        elif line.startswith("    "):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line.strip())
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            document.add_paragraph(line)

    document.save(WORD_OUTPUT)
    return True


def build_pdf(lines):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], alignment=TA_CENTER, fontSize=20, leading=24, spaceAfter=16))
    styles.add(ParagraphStyle(name="CodeBlock", parent=styles["Code"], fontName="Courier", fontSize=8, leading=10, leftIndent=12, rightIndent=12))

    story = []
    in_code = False
    code_lines = []

    for line in lines:
        if line.startswith("```"):
            if in_code:
                story.append(Preformatted("\n".join(code_lines), styles["CodeBlock"]))
                story.append(Spacer(1, 8))
                code_lines = []
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line.strip():
            story.append(Spacer(1, 5))
        elif line.startswith("# "):
            story.append(Paragraph(line[2:], styles["ReportTitle"]))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], styles["Heading1"]))
        elif line.startswith("### "):
            story.append(Paragraph(line[4:], styles["Heading2"]))
        elif line.startswith("- "):
            story.append(Paragraph("&bull; " + line[2:], styles["BodyText"]))
        else:
            safe_line = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe_line.replace("`", ""), styles["BodyText"]))

    document = SimpleDocTemplate(
        str(PDF_OUTPUT),
        pagesize=letter,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
        title="Secure File Encryption Tool - Complete Code Explanation",
    )
    document.build(story)


if __name__ == "__main__":
    source_lines = markdown_lines()
    word_created = build_word(source_lines)
    build_pdf(source_lines)
    if word_created:
        print(WORD_OUTPUT)
    else:
        print("Word export skipped: python-docx is unavailable to this interpreter.")
    print(PDF_OUTPUT)
